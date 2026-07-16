# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = r"c:\Users\PCPV\source\repos\tbs\SMF\docs\Huong_dan_Tarif_Detail.docx"


def set_cell_shading(cell, color_hex):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return h


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.bold = bold
    run.italic = italic
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_shading(hdr_cells[i], "D9E2F3")
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.8)

    title = doc.add_heading("HƯỚNG DẪN SỬ DỤNG", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_heading("Chức năng Tarif Detail (frmTarifDetailInput & frmTarifDetailEdit)", 1)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(doc, "Tài liệu hướng dẫn quản lý chi tiết phí Tarif trong hệ thống SMF.")
    doc.add_paragraph()

    add_heading(doc, "1. Tổng quan", 2)
    add_para(doc, "Hai form frmTarifDetailInput và frmTarifDetailEdit dùng để quản lý chi tiết phí (Detail) trong bảng Deatail_Tarif, gắn với Header Tarif (Header_Tarif). Dữ liệu này được sử dụng khi bấm nút Get trên màn hình Inbound (Debit/Credit) để tạo phí tự động.")

    add_heading(doc, "2. Truy cập chức năng", 2)
    add_table(doc, ["Bước", "Thao tác"], [
        ["1", "Vào menu Report → 19. Tarif Header"],
        ["2", "Mở form Tarif Header (lưới trên: Header, lưới dưới: Detail)"],
        ["3", "Click chọn một Header ở lưới trên"],
        ["4", "Chọn menu tương ứng ở thanh menu"],
    ])
    add_table(doc, ["Menu", "Form mở ra", "Mục đích"], [
        ["New Details", "frmTarifDetailInput", "Thêm mới nhiều dòng chi tiết"],
        ["Edit Detail", "frmTarifDetailEdit", "Sửa một dòng chi tiết đã có"],
        ["Double-click dòng Detail", "frmTarifDetailEdit", "Sửa nhanh"],
    ])
    add_para(doc, "Lưu ý: Phải chọn Header trước. Nếu chưa chọn Header, menu New Details sẽ bị khóa.", italic=True)

    add_heading(doc, "3. frmTarifDetailInput — Thêm chi tiết Tarif mới", 2)
    add_para(doc, "Tiêu đề form: New Tarif Details")
    add_para(doc, "Thanh trên hiển thị: Header: [Tên] | Customer: [Khách hàng] | Type: Debit/Credit")
    add_para(doc, "Quy trình: Nhập thông tin 1 dòng → + Add to List → Lặp lại → Save All", bold=True)
    add_table(doc, ["Bước", "Thao tác"], [
        ["1", "Nhập thông tin từng dòng phí ở vùng Enter Detail Row"],
        ["2", "Bấm + Add to List — dòng được thêm vào lưới Pending Details"],
        ["3", "Lặp lại để thêm nhiều dòng"],
        ["4", "Bấm Save All để ghi tất cả vào database"],
        ["5", "Hoặc Close để thoát (có cảnh báo nếu còn dòng chưa lưu)"],
    ])
    add_table(doc, ["Nút", "Chức năng"], [
        ["+ Add to List", "Thêm dòng hiện tại vào lưới tạm, xóa trắng form để nhập dòng tiếp"],
        ["Remove Row", "Xóa dòng đang chọn khỏi lưới tạm (chưa lưu DB)"],
        ["Save All", "Lưu toàn bộ lưới tạm vào Deatail_Tarif"],
        ["Close", "Đóng form; hỏi xác nhận nếu còn dòng chưa lưu"],
    ])

    add_heading(doc, "4. frmTarifDetailEdit — Sửa chi tiết Tarif", 2)
    add_para(doc, "Tiêu đề form: Edit Tarif Detail")
    add_table(doc, ["Bước", "Thao tác"], [
        ["1", "Form tự động load dữ liệu dòng đang chọn"],
        ["2", "Chỉnh sửa các trường cần thiết"],
        ["3", "Bấm Save để cập nhật database"],
        ["4", "Hoặc Cancel để hủy, không lưu"],
    ])

    add_heading(doc, "5. Mô tả các trường nhập liệu", 2)
    add_table(doc, ["Trường trên form", "Cột DB", "Bắt buộc", "Ghi chú"], [
        ["Item", "itemid", "Có", "Chọn mã phí từ danh mục Charge (Mã/ĐVT/Tên phí)"],
        ["Currency", "currency", "Có", "Loại tiền (USD, VND, …)"],
        ["Unit", "unit", "Không", "Đơn vị tính (BL, KGS, SET, …)"],
        ["Qty", "qty", "Không", "Số lượng"],
        ["Unit Price", "unitprice", "Không", "Đơn giá trước thuế"],
        ["Total Amount", "totalamount", "Tự động", "= Qty × Unit Price"],
        ["VAT%", "vat", "Không", "Phần trăm thuế VAT"],
        ["Unit Price Inc VAT", "unitprice_incvat", "Tự động", "= Unit Price × (1 + VAT% / 100)"],
        ["Tigia", "tigia", "Không", "Tỉ giá quy đổi (quan trọng khi dùng Get trên Inbound)"],
    ])

    add_heading(doc, "6. Tính toán tự động", 2)
    add_para(doc, "Khi thay đổi Qty, Unit Price hoặc VAT%, hệ thống tự động tính:")
    add_para(doc, "• Total Amount = Qty × Unit Price")
    add_para(doc, "• Unit Price Inc VAT = Unit Price × (1 + VAT% / 100)")
    add_para(doc, "Hai trường này chỉ đọc (màu xám), không nhập tay.", italic=True)

    add_heading(doc, "7. Kiểm tra dữ liệu", 2)
    add_para(doc, "• Item phải chọn hợp lệ từ danh sách")
    add_para(doc, "• Currency không được để trống")
    add_para(doc, "• Khi Save All (form Input): phải có ít nhất 1 dòng trong lưới tạm")

    add_heading(doc, "8. Cấu trúc dữ liệu lưu trữ", 2)
    add_para(doc, "Header (Header_Tarif) — tạo ở form Tarif Header:", bold=True)
    add_table(doc, ["Cột", "Ý nghĩa"], [
        ["Name", "Tên bộ Tarif"],
        ["customer_id", "Khách hàng"],
        ["Type", "Debit hoặc Credit"],
        ["usercreate, date", "Người tạo, ngày tạo"],
    ])
    add_para(doc, "Detail (Deatail_Tarif) — tạo/sửa ở 2 form này:", bold=True)
    add_table(doc, ["Cột", "Nguồn"], [
        ["id_header", "Header đang chọn"],
        ["itemid", "Item"],
        ["currency, unit, qty", "Nhập trên form"],
        ["unitprice, totalamount, vat, unitprice_incvat, tigia", "Nhập / tự tính"],
    ])

    add_heading(doc, "9. Liên kết với Inbound (Debit / Credit)", 2)
    add_para(doc, "Sau khi tạo xong Header + Detail:")
    add_para(doc, "1. Vào Inbound (Sea hoặc Air Import)")
    add_para(doc, "2. Chọn job inbound")
    add_para(doc, "3. Tab Debit hoặc Credit → chọn Header Tarif tương ứng loại (Debit / Credit)")
    add_para(doc, "4. Bấm Get (Button116 = Debit, Button117 = Credit)")
    add_para(doc, "Mapping dữ liệu khi import:", bold=True)
    add_table(doc, ["Tarif", "Inbound Freight"], [
        ["customer_id (Header)", "customerid"],
        ["itemid", "itemid"],
        ["currency", "currency"],
        ["unit", "containertype"],
        ["qty", "quantity"],
        ["unitprice", "unitprice_"],
        ["unitprice_incvat", "unitprice"],
        ["totalamount", "price_, pricetruocthue"],
        ["vat", "taxprice"],
        ["tigia", "tigia"],
    ])

    add_heading(doc, "10. Ví dụ thực tế", 2)
    add_para(doc, "Tạo Tarif Debit cho khách ABC — phí vận chuyển + phí xếp dỡ:")
    add_para(doc, "1. Tarif Header → New Header: Name = ABC - Debit chuẩn, Customer = ABC Logistics, Type = Debit → Save")
    add_para(doc, "2. Chọn Yes khi hỏi Do you want to add details now?")
    add_para(doc, "3. Dòng 1: Item = FRT/KGS/Freight, Currency = USD, Unit = KGS, Qty = 100, Unit Price = 2.5, VAT% = 8, Tigia = 25400 → + Add to List")
    add_para(doc, "4. Dòng 2: Item = THC/BL/Terminal Handling, Currency = USD, Unit = BL, Qty = 1, Unit Price = 150, VAT% = 8, Tigia = 25400 → + Add to List")
    add_para(doc, "5. Save All → 2 dòng lưu vào Deatail_Tarif")
    add_para(doc, "6. Trên Inbound → chọn Tarif ABC - Debit chuẩn → Get → 2 dòng debit được tạo.")

    add_heading(doc, "11. Lưu ý quan trọng", 2)
    add_para(doc, "• Type Header (Debit / Credit) phải khớp với tab sử dụng trên Inbound")
    add_para(doc, "• New Details cho phép thêm nhiều dòng một lần; Edit Detail chỉ sửa một dòng")
    add_para(doc, "• Xóa Header trên frmTarifHeader sẽ xóa toàn bộ Detail liên quan")
    add_para(doc, "• Bấm Refresh trên frmTarifHeader để cập nhật lưới sau khi lưu")
    add_para(doc, "• Dữ liệu Detail chưa Save All (form Input) sẽ mất nếu đóng form")
    add_para(doc, "• Nên nhập đầy đủ Tigia trong Tarif Detail để dữ liệu chính xác khi import sang Inbound")

    add_heading(doc, "12. Tóm tắt nhanh", 2)
    add_table(doc, ["Tình huống", "Form dùng", "Hành động"], [
        ["Thêm nhiều phí mới", "frmTarifDetailInput", "Add to List → Save All"],
        ["Sửa một phí đã có", "frmTarifDetailEdit", "Save"],
        ["Xem danh sách phí", "frmTarifHeader", "Chọn Header → xem lưới Detail"],
        ["Dùng Tarif trên job", "Inbound", "Chọn Header Tarif → Get"],
    ])

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
